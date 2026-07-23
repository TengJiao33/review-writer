from __future__ import annotations

import hashlib
import json
import importlib.util
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

from docx import Document
from pypdf import PdfWriter


REPO = Path(__file__).resolve().parents[1]
MERGE = REPO / "skills" / "review-draft-merge-polish" / "scripts" / "merge_review.py"
AUDIT = REPO / "skills" / "review-final-audit-release" / "scripts" / "final_audit_scan.py"
MD2DOCX = REPO / "skills" / "review-export-docx" / "scripts" / "md2docx.py"
DOCX_AUDIT = REPO / "skills" / "review-export-docx" / "scripts" / "audit_docx.py"
MATRIX_VALIDATOR = REPO / "skills" / "review-literature-matrix-outline" / "scripts" / "validate_evidence_matrix.py"
BLUEPRINT_VALIDATOR = REPO / "skills" / "review-section-blueprint" / "scripts" / "validate_blueprint.py"
BLUEPRINT_INIT = REPO / "skills" / "review-section-blueprint" / "scripts" / "init_section_blueprint.py"
DRAFT_VALIDATOR = REPO / "skills" / "review-section-drafting-figure-picking" / "scripts" / "validate_section_drafts.py"
SCREENING_VALIDATOR = REPO / "skills" / "review-topic-paper-discovery" / "scripts" / "validate_screening.py"
METADATA_PREP = REPO / "skills" / "review-metadata-prep" / "scripts" / "prepare_metadata.py"
DISCOVER = REPO / "skills" / "review-topic-paper-discovery" / "scripts" / "discover.py"
FIGURE_SELECTOR = REPO / "skills" / "review-section-drafting-figure-picking" / "scripts" / "select_initial_figure_candidates.py"
FIGURE_REDRAW = REPO / "skills" / "review-figure-style-redraw" / "scripts" / "redraw_figures.py"
PROJECT_STATUS = REPO / "skills" / "review-writing-orchestrator" / "scripts" / "project_status.py"
RUN_AND_RECORD = REPO / "skills" / "review-writing-orchestrator" / "scripts" / "run_and_record.py"
EXTERNAL_INGEST = REPO / "skills" / "review-topic-paper-discovery" / "scripts" / "ingest_external_papers.py"
MINERU_PARSER = REPO / "skills" / "mineru-precise-parse-review-writer" / "scripts" / "parse_review_writer_pdfs.py"
PORTFOLIO_BUILDER = REPO / "skills" / "review-literature-matrix-outline" / "scripts" / "build_review_portfolio.py"
COMPARISON_TABLE = REPO / "skills" / "review-section-drafting-figure-picking" / "scripts" / "build_method_comparison_table.py"
RENDER_DOCX = REPO / "skills" / "review-export-docx" / "scripts" / "render_docx.py"


def write_json(path: Path, payload) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def run(*args: str) -> subprocess.CompletedProcess[str]:
    return subprocess.run(
        [sys.executable, *map(str, args)],
        capture_output=True,
        text=True,
        encoding="utf-8",
    )


class ReviewWorkflowContractsTest(unittest.TestCase):
    def setUp(self) -> None:
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.project_id = "fixture-review"
        self.project = self.root / "review-projects" / self.project_id
        self._write_matrix_and_metadata()
        write_json(self.project / "02_section_drafting" / "figure_candidates.json", [])
        write_json(
            self.project / "02_section_drafting" / "paper_figure_inventory.json",
            {"candidate_count": 0, "candidates": [], "papers": []},
        )
        write_json(
            self.project / "02_section_drafting" / "paper_figure_candidates.json",
            {"inventory_candidate_count": 0, "candidates": []},
        )
        write_json(
            self.project / "00_discovery" / "screening_validation.json",
            {"blocking_issues": [], "warnings": []},
        )

    def tearDown(self) -> None:
        self.temp.cleanup()

    def test_rule_pack_selection_uses_general_default_and_domain_match(self) -> None:
        spec = importlib.util.spec_from_file_location("blueprint_init", BLUEPRINT_INIT)
        module = importlib.util.module_from_spec(spec)
        assert spec.loader is not None
        spec.loader.exec_module(module)
        skill_root = BLUEPRINT_INIT.parents[1]

        self.assertEqual(
            module.select_rule_pack(skill_root, "enzymatic PET depolymerization")[0],
            "general",
        )
        self.assertEqual(
            module.select_rule_pack(skill_root, "enantioselective allene synthesis from propargylic alcohols")[0],
            "allenation",
        )

    def test_blueprint_ready_status_requires_claim_level_evidence_links(self) -> None:
        path = self.project / "01_matrix_outline" / "section_blueprint.json"
        payload = {
            "status": "ready_for_drafting",
            "coverage_contract": {"dimensions": []},
            "sections": [
                {
                    "section_id": "sec1",
                    "title": "Evidence boundary",
                    "review_claims": [
                        {
                            "claim": "The two methods answer different practical questions.",
                            "claim_type": "comparison",
                            "supporting_papers": ["P001", "P002"],
                            "evidence_strength": "needs verification",
                        }
                    ],
                }
            ],
        }
        write_json(path, payload)
        result = run(BLUEPRINT_VALIDATOR, "--review-root", self.root, "--project-id", self.project_id)
        self.assertNotEqual(result.returncode, 0)
        report = json.loads(
            (self.project / "01_matrix_outline" / "blueprint_validation.json").read_text(
                encoding="utf-8"
            )
        )
        self.assertTrue(any("no linked evidence_ids" in issue for issue in report["blocking_issues"]))
        self.assertTrue(any("unverified evidence strength" in issue for issue in report["blocking_issues"]))

        claim = payload["sections"][0]["review_claims"][0]
        claim["evidence_ids"] = ["P001-E01", "P002-E01"]
        claim["evidence_strength"] = "cross_checked_full_text"
        write_json(path, payload)
        result = run(BLUEPRINT_VALIDATOR, "--review-root", self.root, "--project-id", self.project_id)
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)

    def test_metadata_prep_only_slug_does_not_reprocess_the_library(self) -> None:
        registry = self.root / "review-library" / "registry" / "papers.jsonl"
        registry.parent.mkdir(parents=True, exist_ok=True)
        registry.write_text(
            "".join(
                json.dumps({"paper_id": paper_id, "source_pdf": f"existing-{paper_id}.pdf"}) + "\n"
                for paper_id in ("P001", "P002")
            ),
            encoding="utf-8",
        )
        pdf_root = self.root / "chem_papers"
        mineru_root = self.root / "mineru-outputs"
        slugs = []
        spec = importlib.util.spec_from_file_location("metadata_prep", METADATA_PREP)
        module = importlib.util.module_from_spec(spec)
        assert spec.loader is not None
        spec.loader.exec_module(module)
        for name in ("first-study.pdf", "second-study.pdf"):
            pdf = pdf_root / "web-imports" / name
            pdf.parent.mkdir(parents=True, exist_ok=True)
            writer = PdfWriter()
            writer.add_blank_page(width=612, height=792)
            with pdf.open("wb") as handle:
                writer.write(handle)
            relative_stem = str(pdf.relative_to(pdf_root).with_suffix(""))
            slug = module.slugify_mineru(relative_stem)
            slugs.append(slug)
            markdown = mineru_root / "markdown" / f"{slug}.md"
            markdown.parent.mkdir(parents=True, exist_ok=True)
            markdown.write_text(
                f"# {name.removesuffix('.pdf')}\n\nAbstract\n\nA bounded imported study reports a result.\n",
                encoding="utf-8",
            )
            write_json(
                mineru_root / "extracted" / slug / f"{slug}_content_list.json",
                [{"type": "text", "text": f"{name} imported study", "page_idx": 0}],
            )

        result = run(
            METADATA_PREP,
            "--review-root",
            self.root,
            "--mineru-output",
            mineru_root,
            "--pdf-root",
            pdf_root,
            "--discover-from-pdf-root",
            "--append-registry",
            "--only-slug",
            slugs[0],
        )
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        self.assertTrue((self.root / "review-library" / "metadata" / "papers" / "P003.metadata.json").exists())
        self.assertFalse((self.root / "review-library" / "metadata" / "papers" / "P004.metadata.json").exists())
        registry_rows = [json.loads(line) for line in registry.read_text(encoding="utf-8").splitlines()]
        self.assertEqual([row["paper_id"] for row in registry_rows], ["P001", "P002", "P003"])

    def test_generic_method_cards_produce_a_traceable_comparison_table(self) -> None:
        matrix_path = self.project / "01_matrix_outline" / "literature_matrix.json"
        payload = json.loads(matrix_path.read_text(encoding="utf-8"))
        fields = {
            "study_design": "controlled primary study",
            "subject_or_substrate": "propargylic carbonate substrates",
            "intervention_or_method": "catalytic carbonylation",
            "conditions_or_context": "reported reaction conditions",
            "outcome_or_metric": "isolated product yield",
            "main_result": "representative products were obtained",
            "limitations": "scope remains bounded",
        }
        for row in payload["papers"]:
            evidence_id = f"{row['paper_id']}-E01"
            row["method_card"] = {
                **fields,
                "method_field_evidence": {
                    field: [evidence_id]
                    for field in fields
                },
            }
        write_json(matrix_path, payload)
        write_json(
            self.project / "00_discovery" / "topic_contract.json",
            {
                "topic": "Comparative catalytic methods",
                "review_profile": "focused",
                "central_question": "Which method is useful under which conditions?",
                "important_coverage": ["method choice"],
            },
        )
        write_json(
            self.project / "00_discovery" / "selected_discovery_results.json",
            {
                "screening_decisions": [
                    {
                        "paper_id": row["paper_id"],
                        "decision": "include",
                        "portfolio_intent_hint": "core",
                    }
                    for row in payload["papers"]
                ]
            },
        )
        result = run(PORTFOLIO_BUILDER, "--review-root", self.root, "--project-id", self.project_id)
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        cards = json.loads(
            (self.project / "01_matrix_outline" / "method_cards.json").read_text(encoding="utf-8")
        )["method_cards"]
        self.assertTrue(all(card["recording_status"] == "recorded_with_field_provenance" for card in cards))
        self.assertTrue(all("outcome_or_metric" in card["populated_fields"] for card in cards))

        result = run(
            COMPARISON_TABLE,
            "--review-root",
            self.root,
            "--project-id",
            self.project_id,
            "--require-traceable",
        )
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        table = (
            self.project / "02_section_drafting" / "method_comparison_table.md"
        ).read_text(encoding="utf-8")
        self.assertIn("Outcome / metric", table)
        self.assertIn("catalytic carbonylation", table)

    def _write_matrix_and_metadata(self) -> None:
        rows = []
        for paper_id, title, year in (
            ("P001", "Palladium Carbonylation of Propargylic Carbonates", 2020),
            ("P002", "Nickel Electrocarboxylation of Propargylic Carbonates", 2022),
        ):
            rows.append(
                {
                    "paper_id": paper_id,
                    "title": title,
                    "authors": ["Ada Chemist", "Ben Researcher"],
                    "keywords": {"substrate": "propargylic carbonates"},
                    "abstract": "Evidence-grounded abstract.",
                    "main_content": "Evidence " * 360,
                    "most_relevant_figure": {},
                    "evidence_anchors": [
                        {
                            "evidence_id": f"{paper_id}-E01",
                            "note": "The full text reports the transformation and representative scope.",
                            "source_excerpt": "The transformation and representative substrate scope are reported with quantified yields and operating conditions.",
                            "source_path": f"review-library/sources/{paper_id}.md",
                            "locator": "Results, paragraph 1",
                            "source_level": "full_text",
                            "evidence_kind": "result",
                            "certainty": "direct",
                        },
                        {
                            "evidence_id": f"{paper_id}-E02",
                            "note": "The authors describe the pathway as a mechanistic proposal.",
                            "source_excerpt": "The pathway is presented by the authors as a mechanistic proposal based on intermediate trapping evidence.",
                            "source_path": f"review-library/sources/{paper_id}.md",
                            "locator": "Mechanism, paragraph 1",
                            "source_level": "full_text",
                            "evidence_kind": "mechanism",
                            "certainty": "author_interpretation",
                        },
                    ],
                }
            )
            source = self.root / "review-library" / "sources" / f"{paper_id}.md"
            source.parent.mkdir(parents=True, exist_ok=True)
            source.write_text(
                "# Results\n\nThe transformation and representative substrate scope are reported with quantified yields and operating conditions.\n\n"
                "# Mechanism\n\nThe pathway is presented by the authors as a mechanistic proposal based on intermediate trapping evidence.\n",
                encoding="utf-8",
            )
            write_json(
                self.root / "review-library" / "metadata" / "papers" / f"{paper_id}.metadata.json",
                {
                    "paper_id": paper_id,
                    "title": {"value": title},
                    "authors": {"value": ["Ada Chemist", "Ben Researcher"]},
                    "journal": {"value": "Journal of Allene Chemistry"},
                    "year": {"value": year},
                    "doi": {"value": f"10.1000/{paper_id.lower()}"},
                },
            )
        write_json(self.project / "01_matrix_outline" / "literature_matrix.json", {"papers": rows})

    def _front_matter(self) -> dict:
        abstract = (
            "This review compares catalytic strategies for converting propargylic carbonates into allenes. "
            "It evaluates activation mode, selectivity control, mechanistic evidence, substrate scope, and practical limitations. "
        ) * 12
        return {
            "title": "Catalytic Allene Synthesis from Propargylic Carbonates",
            "abstract": abstract,
            "keywords": ["allenes", "propargylic carbonates", "palladium", "nickel", "selectivity"],
        }

    def _source_receipts(self, *evidence_ids: str) -> list[dict]:
        receipts = []
        for evidence_id in evidence_ids:
            paper_id = evidence_id.split("-", 1)[0]
            source = self.root / "review-library" / "sources" / f"{paper_id}.md"
            excerpt = (
                "The pathway is presented by the authors as a mechanistic proposal based on intermediate trapping evidence."
                if evidence_id.endswith("E02")
                else "The transformation and representative substrate scope are reported with quantified yields and operating conditions."
            )
            receipts.append(
                {
                    "paper_id": paper_id,
                    "evidence_id": evidence_id,
                    "source_path": str(source),
                    "source_sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
                    "locator": "Mechanism, paragraph 1" if evidence_id.endswith("E02") else "Results, paragraph 1",
                    "checked_excerpt": excerpt,
                }
            )
        return receipts

    def test_legacy_numeric_citations_are_rejected(self) -> None:
        payload = {
            "front_matter": self._front_matter(),
            "sections": [
                {
                    "section_id": "sec1",
                    "title": "1. Introduction",
                    "paragraphs": [
                        {
                            "paragraph_id": "sec1-p1",
                            "paragraph_type": "comparison",
                            "markdown": "Palladium and nickel systems differ [1].",
                            "cited_paper_ids": ["P001", "P002"],
                        }
                    ],
                }
            ],
        }
        write_json(self.project / "02_section_drafting" / "section_drafts.json", payload)
        result = run(MERGE, "--review-root", self.root, "--project-id", self.project_id)
        self.assertNotEqual(result.returncode, 0)
        self.assertIn("numeric citations are forbidden", result.stdout)

    def test_abstract_image_embed_is_rejected(self) -> None:
        front = self._front_matter()
        front["abstract"] += "\n\n![A misplaced scheme](figures/scheme.png)"
        write_json(
            self.project / "02_section_drafting" / "section_drafts.json",
            {
                "front_matter": front,
                "sections": [
                    {
                        "section_id": "sec1",
                        "title": "1. Introduction",
                        "paragraphs": [
                            {
                                "paragraph_id": "sec1-p1",
                                "markdown": "The focal transformation is reported [@P001].",
                                "evidence_ids": ["P001-E01"],
                            }
                        ],
                    }
                ],
            },
        )
        result = run(MERGE, "--review-root", self.root, "--project-id", self.project_id)
        self.assertNotEqual(result.returncode, 0)
        self.assertIn("front_matter.abstract must not contain embedded images", result.stdout)

    def test_delegated_screening_uses_generic_topic_contract(self) -> None:
        write_json(
            self.project / "00_discovery" / "topic_contract.json",
            {
                "topic": "Interventions and long-term retention",
                "review_profile": "focused",
                "central_question": "Which interventions improve long-term retention?",
                "important_coverage": ["follow-up duration", "population"],
                "inclusion_criteria": ["reports a retention outcome"],
                "exclusion_criteria": ["no longitudinal measurement"],
            },
        )
        write_json(
            self.project / "00_discovery" / "selected_discovery_results.json",
            {
                "screening": {"status": "confirmed", "decided_by": "agent"},
                "candidate_paper_ids": ["P001", "P002"],
                "local_papers": [{"paper_id": "P001"}],
                "screening_decisions": [
                    {
                        "paper_id": "P001",
                        "decision": "include",
                        "study_type": "primary_research",
                        "relevance_summary": "Reports the outcome named in the central question.",
                        "decision_basis": "Full-text results and the inclusion criterion.",
                    },
                    {
                        "paper_id": "P002",
                        "decision": "exclude",
                        "relevance_summary": "Measures only immediate performance.",
                        "decision_basis": "No longitudinal measurement.",
                    },
                ],
            },
        )
        result = run(
            SCREENING_VALIDATOR,
            "--review-root",
            self.root,
            "--project-id",
            self.project_id,
        )
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)

    def test_journal_hint_does_not_match_affiliation_text(self) -> None:
        spec = importlib.util.spec_from_file_location("review_metadata_prep", METADATA_PREP)
        self.assertIsNotNone(spec)
        self.assertIsNotNone(spec.loader)
        module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(module)
        extracted = module.extract_journal(
            "Shanghai Key Laboratory of Green Chemistry and Chemical Processes",
            "article.pdf",
        )
        self.assertIsNone(extracted["value"])
        filename_extracted = module.extract_journal("", "Green Chemistry - article.pdf")
        self.assertEqual(filename_extracted["value"], "Green Chemistry")
        abbreviated = module.extract_journal(
            "",
            "46-Angew Chem Int Ed - 2012 - Wan - Enantioselective Amination.pdf",
        )
        self.assertEqual(abbreviated["value"], "Angew. Chem. Int. Ed.")
        regenerated = {
            "title": {"value": "Parsed title", "source": "markdown", "human_checked": False},
            "doi": {"value": None, "source": "rule_not_found", "human_checked": False},
            "authors": {"value": ["Parsed Author"], "source": "markdown", "human_checked": False},
        }
        existing = {
            "title": {"value": "Checked title", "source": "manual", "human_checked": True},
            "doi": {
                "value": "10.1000/persisted",
                "source": "external_discovery",
                "human_checked": False,
            },
            "authors": {"value": ["Old Guess"], "source": "markdown", "human_checked": False},
        }
        module.preserve_human_checked_fields(regenerated, existing)
        self.assertEqual(regenerated["title"]["value"], "Checked title")
        self.assertEqual(regenerated["doi"]["value"], "10.1000/persisted")
        self.assertEqual(regenerated["authors"]["value"], ["Parsed Author"])
        orphan_markdown = self.root / "mineru-outputs" / "markdown" / "web-imports-orphan.md"
        orphan_markdown.parent.mkdir(parents=True, exist_ok=True)
        orphan_markdown.write_text("# Repository article\n", encoding="utf-8")
        orphan_xml = self.root / "chem_papers" / "web-imports" / "orphan.jats.xml"
        orphan_xml.parent.mkdir(parents=True, exist_ok=True)
        orphan_xml.write_text("<article/>", encoding="utf-8")
        jobs = module.jobs_from_pdf_root(
            self.root / "chem_papers",
            self.root / "mineru-outputs",
        )
        self.assertEqual(len(jobs), 1)
        self.assertEqual(jobs[0]["source_xml"], str(orphan_xml))
        constrained = module.constrain_structured_tags(
            {
                "ligand_or_chiral_source": "enantioselective synthesis",
                "document_scope": "primary research article",
            },
            {key: ["not specified"] for key in module.STRUCTURED_TAG_KEYS},
        )
        self.assertTrue(all(value == "not specified" for value in constrained.values()))

    def test_discovery_selection_keeps_compact_screening_context(self) -> None:
        sys.path.insert(0, str(DISCOVER.parent))
        try:
            spec = importlib.util.spec_from_file_location("review_discover", DISCOVER)
            self.assertIsNotNone(spec)
            self.assertIsNotNone(spec.loader)
            module = importlib.util.module_from_spec(spec)
            spec.loader.exec_module(module)
        finally:
            sys.path.pop(0)
        selected = module.selected_from_combined(
            [
                {
                    "keyword": "nickel catalysis",
                    "category": "catalyst_or_method",
                    "keep": True,
                    "local_results": [
                        {
                            "paper_id": "P001",
                            "title": "A paper",
                            "year": 2025,
                            "journal": "A Journal",
                            "abstract": "Enough context for screening.",
                            "structured_tags": {"catalyst_or_method": "nickel catalysis"},
                            "source_paths": {"markdown": "paper.md", "pdf": "paper.pdf"},
                            "role": "core_candidate",
                            "score": 0.8,
                            "keep": True,
                        }
                    ],
                    "web_results": [],
                }
            ]
        )
        paper = selected["local_papers"][0]
        self.assertEqual(paper["abstract"], "Enough context for screening.")
        self.assertEqual(paper["structured_tags"]["catalyst_or_method"], "nickel catalysis")
        self.assertEqual(paper["source_paths"]["markdown"], "paper.md")

    def test_discovery_builds_generic_contract_queries_outside_allene_domain(self) -> None:
        sys.path.insert(0, str(DISCOVER.parent))
        try:
            spec = importlib.util.spec_from_file_location("review_discover_fallback", DISCOVER)
            self.assertIsNotNone(spec)
            self.assertIsNotNone(spec.loader)
            module = importlib.util.module_from_spec(spec)
            spec.loader.exec_module(module)
        finally:
            sys.path.pop(0)
        topic = "Covalent organic frameworks for photocatalytic carbon dioxide reduction"
        keyword_set = module.build_keyword_set(
            topic,
            [],
            {
                "manuscript_title": topic,
                "central_question": "Which framework structures control activity and selectivity?",
                "important_coverage": [
                    "Linker and node chemistry in carbon dioxide reduction.",
                    "Operando evidence for charge transfer and reaction intermediates.",
                ],
            },
        )
        self.assertGreaterEqual(len(keyword_set["agent_keywords"]), 3)
        self.assertEqual(
            keyword_set["merged_keywords"][0]["category"],
            "core_topic",
        )
        self.assertNotIn(
            "topic_fallback",
            {
                source
                for row in keyword_set["merged_keywords"]
                for source in row["source"]
            },
        )

    def test_sulfide_topic_gets_generic_queries_not_allene_categories(self) -> None:
        sys.path.insert(0, str(DISCOVER.parent))
        try:
            spec = importlib.util.spec_from_file_location("review_discover_domain", DISCOVER)
            self.assertIsNotNone(spec)
            self.assertIsNotNone(spec.loader)
            module = importlib.util.module_from_spec(spec)
            spec.loader.exec_module(module)
        finally:
            sys.path.pop(0)
        solid = module.infer_keywords(
            "Interphase chemistry at sulfide solid-electrolyte and lithium-metal interfaces",
            [],
            {
                "important_coverage": [
                    "Argyrodite and LGPS cathode interfaces and contact loss.",
                    "Lithium-metal interphases, voiding, dendrites, and pressure.",
                ]
            },
        )
        self.assertGreaterEqual(len(solid), 3)
        self.assertTrue(
            all(
                row["category"]
                in {"core_topic", "coverage", "mechanism_or_outcome", "scope"}
                for row in solid
            )
        )
        self.assertEqual(
            module.classify_keyword("sulfide solid electrolyte interface"),
            "user_query",
        )

    def test_markdown_topic_contract_accepts_common_headings_and_continuations(self) -> None:
        sys.path.insert(0, str(DISCOVER.parent))
        try:
            spec = importlib.util.spec_from_file_location("review_discover_contract", DISCOVER)
            self.assertIsNotNone(spec)
            self.assertIsNotNone(spec.loader)
            module = importlib.util.module_from_spec(spec)
            spec.loader.exec_module(module)
        finally:
            sys.path.pop(0)
        path = self.root / "topic_input.md"
        path.write_text(
            "\n".join(
                [
                    "# Review Topic",
                    "## Manuscript Title",
                    "Interfaces in sulfide batteries",
                    "## Important Coverage",
                    "- Cathode interfaces, including coatings,",
                    "  contact loss, and transport.",
                    "- Lithium-metal interphases.",
                    "## Inclusion",
                    "Include primary interface studies with full text.",
                    "## Exclusion",
                    "Exclude liquid-only systems.",
                ]
            ),
            encoding="utf-8",
        )
        contract = module.load_topic_contract_file(str(path))
        self.assertEqual(
            contract["important_coverage"][0],
            "Cathode interfaces, including coatings, contact loss, and transport.",
        )
        self.assertEqual(
            contract["inclusion_criteria"],
            ["Include primary interface studies with full text."],
        )
        self.assertEqual(
            contract["exclusion_criteria"],
            ["Exclude liquid-only systems."],
        )

    def test_generic_local_retrieval_does_not_require_chemistry_tags(self) -> None:
        sys.path.insert(0, str(DISCOVER.parent))
        try:
            spec = importlib.util.spec_from_file_location("review_discover_local", DISCOVER)
            self.assertIsNotNone(spec)
            self.assertIsNotNone(spec.loader)
            module = importlib.util.module_from_spec(spec)
            spec.loader.exec_module(module)
        finally:
            sys.path.pop(0)
        papers = {
            "P900": {
                "paper_id": "P900",
                "title": {"value": "Operando imaging of lithium metal void formation"},
                "abstract": {
                    "value": "Pressure-dependent contact loss at a solid electrolyte interface."
                },
                "structured_tags": {
                    key: "not specified" for key in module.STRUCTURED_TAG_KEYS
                },
                "source_paths": {},
            }
        }
        grouped = module.local_search_by_keyword(
            papers,
            [
                {
                    "keyword": "lithium metal void formation",
                    "category": "mechanism_or_outcome",
                    "keep": True,
                }
            ],
            "solid-state battery interfaces",
            {key: {} for key in module.STRUCTURED_TAG_KEYS},
        )
        self.assertEqual(
            [row["paper_id"] for row in grouped[0]["local_results"]],
            ["P900"],
        )
        self.assertIn("title", grouped[0]["local_results"][0]["matched_fields"])

    def test_external_ingest_rejects_in_progress_or_mixed_discovery_runs(self) -> None:
        discovery = (
            self.root
            / "review-projects"
            / self.project_id
            / "00_discovery"
        )
        discovery.mkdir(parents=True, exist_ok=True)
        write_json(
            discovery / ".discovery_in_progress.json",
            {"discovery_run_id": "run-active"},
        )
        active = run(
            EXTERNAL_INGEST,
            "--review-root",
            self.root,
            "--project-id",
            self.project_id,
        )
        self.assertNotEqual(active.returncode, 0)
        self.assertIn("Discovery is still in progress", active.stdout + active.stderr)

        (discovery / ".discovery_in_progress.json").unlink()
        write_json(
            discovery / "external_ingest_plan.json",
            {"discovery_run_id": "run-plan", "items": []},
        )
        write_json(
            discovery / "selected_discovery_results.json",
            {"discovery_run_id": "run-selected"},
        )
        mixed = run(
            EXTERNAL_INGEST,
            "--review-root",
            self.root,
            "--project-id",
            self.project_id,
        )
        self.assertNotEqual(mixed.returncode, 0)
        self.assertIn(
            "Discovery outputs come from different runs",
            mixed.stdout + mixed.stderr,
        )

    def test_crossref_query_and_editorial_filter_are_unbiased(self) -> None:
        sys.path.insert(0, str(DISCOVER.parent))
        try:
            spec = importlib.util.spec_from_file_location("review_discover_crossref", DISCOVER)
            self.assertIsNotNone(spec)
            self.assertIsNotNone(spec.loader)
            module = importlib.util.module_from_spec(spec)
            spec.loader.exec_module(module)
        finally:
            sys.path.pop(0)
        with mock.patch.object(
            module,
            "crossref_search_items",
            return_value=[
                {
                    "DOI": "10.1039/example/v1/review1",
                    "title": ["Review for an otherwise relevant paper"],
                    "type": "peer-review",
                },
                {
                    "DOI": "10.1000/article",
                    "title": ["PET enzymatic depolymerization"],
                    "type": "journal-article",
                    "issued": {"date-parts": [[2024]]},
                },
            ],
        ) as search:
            rows = module.web_search("PET enzymatic depolymerization", "PET recycling", 8)
        query = search.call_args.args[0]
        self.assertNotIn("review paper DOI", query)
        self.assertEqual(query, "PET enzymatic depolymerization")
        self.assertEqual([row["doi"] for row in rows], ["10.1000/article"])
        self.assertNotIn("example@example.com", module.CROSSREF_USER_AGENT)

        captured: list[str] = []

        def fake_request(url: str):
            captured.append(url)
            return {"message": {"items": []}}

        with mock.patch.object(module, "crossref_request_json", side_effect=fake_request):
            module.crossref_search_items("PET depolymerization", 20)
        self.assertIn("filter=type%3Ajournal-article", captured[0])

    def test_crossref_reference_expansion_recovers_and_locates_cited_work(self) -> None:
        sys.path.insert(0, str(DISCOVER.parent))
        try:
            spec = importlib.util.spec_from_file_location("review_discover_expansion", DISCOVER)
            self.assertIsNotNone(spec)
            self.assertIsNotNone(spec.loader)
            module = importlib.util.module_from_spec(spec)
            spec.loader.exec_module(module)
        finally:
            sys.path.pop(0)
        seed = {
            "DOI": "10.1000/review",
            "title": ["A review of enzymatic PET depolymerization"],
            "type": "journal-article",
            "issued": {"date-parts": [[2023]]},
            "reference": [
                {
                    "DOI": "10.1126/science.aad6359",
                    "article-title": "A bacterium that degrades and assimilates poly(ethylene terephthalate)",
                    "year": "2016",
                }
            ],
        }
        cited = {
            "DOI": "10.1126/science.aad6359",
            "title": ["A bacterium that degrades and assimilates poly(ethylene terephthalate)"],
            "type": "journal-article",
            "issued": {"date-parts": [[2016]]},
            "license": [{"URL": "https://creativecommons.org/licenses/by/4.0/"}],
            "link": [{"URL": "https://example.org/pet.pdf", "content-type": "application/pdf"}],
        }

        def fake_work(doi: str):
            return seed if "review" in doi else cited

        with mock.patch.object(module, "crossref_search_items", return_value=[seed]), mock.patch.object(
            module, "fetch_crossref_work", side_effect=fake_work
        ), mock.patch.object(module, "fetch_crossref_works", return_value=[cited]):
            expansion = module.crossref_reference_expansion(
                "Enzymatic depolymerization of polyethylene terephthalate for chemical recycling",
                ["PET enzymatic depolymerization"],
                seed_limit=1,
                result_limit=5,
            )
        self.assertEqual(expansion["status"], "ok")
        recovered = next(
            row for row in expansion["results"] if row.get("doi") == "10.1126/science.aad6359"
        )
        self.assertEqual(recovered["source"], "crossref_reference_expansion")
        self.assertEqual(recovered["open_access_pdf_url"], "https://example.org/pet.pdf")

    def test_discovery_reads_structured_markdown_topic_contract(self) -> None:
        sys.path.insert(0, str(DISCOVER.parent))
        try:
            spec = importlib.util.spec_from_file_location("review_discover_contract", DISCOVER)
            self.assertIsNotNone(spec)
            self.assertIsNotNone(spec.loader)
            module = importlib.util.module_from_spec(spec)
            spec.loader.exec_module(module)
        finally:
            sys.path.pop(0)
        contract = self.root / "topic_input.md"
        contract.write_text(
            "# Review Topic\n\n"
            "## Manuscript Title\n\nA concise title\n\n"
            "## Retrieval Query\n\nNickel catalysis with propargylic electrophiles\n\n"
            "## Central Question\n\nWhich factors control product topology?\n\n"
            "## Important Coverage\n\n- activation mode\n- ligand effects\n\n"
            "## Inclusion Criteria\n\n- nickel is part of the transformation\n\n"
            "## Exclusion Criteria\n\n- nickel is not part of the transformation\n\n"
            "## Suggested Retrieval Keywords\n\n- nickel catalysis\n- propargylic electrophiles\n",
            encoding="utf-8",
        )
        parsed = module.load_topic_contract_file(str(contract))
        self.assertEqual(parsed["manuscript_title"], "A concise title")
        self.assertEqual(parsed["retrieval_query"], "Nickel catalysis with propargylic electrophiles")
        self.assertEqual(parsed["important_coverage"], ["activation mode", "ligand effects"])
        self.assertEqual(parsed["suggested_keywords"], ["nickel catalysis", "propargylic electrophiles"])

    def test_provider_status_distinguishes_irrelevance_from_failure(self) -> None:
        sys.path.insert(0, str(DISCOVER.parent))
        try:
            spec = importlib.util.spec_from_file_location("review_discover_provider", DISCOVER)
            self.assertIsNotNone(spec)
            self.assertIsNotNone(spec.loader)
            module = importlib.util.module_from_spec(spec)
            spec.loader.exec_module(module)
        finally:
            sys.path.pop(0)
        records, errors = module.split_provider_rows(
            [
                {
                    "title": "A valid but irrelevant record",
                    "keep": False,
                    "reason": "low topic overlap",
                },
                {
                    "title": "WEB_SEARCH_FAILED: TimeoutError",
                    "keep": False,
                    "reason": "timed out",
                },
            ]
        )
        self.assertEqual([row["title"] for row in records], ["A valid but irrelevant record"])
        self.assertEqual(errors, ["timed out"])
        self.assertEqual(
            module.provider_run_status(
                True,
                {
                    "attempted_queries": 1,
                    "successful_queries": 1,
                    "returned_records": 1,
                    "retained_records": 0,
                },
                [],
            ),
            "no_retained_results",
        )

    def test_external_discovery_builds_deduplicated_mineru_ingest_plan(self) -> None:
        sys.path.insert(0, str(DISCOVER.parent))
        try:
            spec = importlib.util.spec_from_file_location("review_discover_ingest", DISCOVER)
            self.assertIsNotNone(spec)
            self.assertIsNotNone(spec.loader)
            module = importlib.util.module_from_spec(spec)
            spec.loader.exec_module(module)
        finally:
            sys.path.pop(0)
        local = {
            "P001": {
                "paper_id": "P001",
                "title": {"value": "Known Allene Paper"},
                "doi": {"value": "10.1000/known"},
            }
        }
        rows = module.annotate_external_results(
            [
                {
                    "external_id": "s2-known",
                    "title": "Known Allene Paper",
                    "doi": "https://doi.org/10.1000/KNOWN",
                    "source": "semantic_scholar",
                    "keep": True,
                    "score": 0.8,
                },
                {
                    "external_id": "s2-new",
                    "title": "New Open Allene Paper",
                    "authors": ["Ada Author"],
                    "doi": "10.1000/new",
                    "year": 2026,
                    "journal": "Open Chemistry",
                    "abstract": "A discovery abstract retained for metadata reconciliation.",
                    "open_access_pdf_url": "https://example.org/new.pdf",
                    "source": "semantic_scholar",
                    "keep": True,
                    "score": 0.7,
                },
            ],
            local,
        )
        selected = module.selected_from_combined(
            [
                {
                    "keyword": "allene synthesis",
                    "category": "reaction_type",
                    "keep": True,
                    "local_results": [],
                    "web_results": rows,
                },
                {
                    "keyword": "nickel catalysis",
                    "category": "catalyst_or_method",
                    "keep": True,
                    "local_results": [],
                    "web_results": [rows[1]],
                },
            ]
        )
        self.assertEqual(len(selected["web_papers"]), 2)
        new_row = next(row for row in selected["web_papers"] if row["external_id"] == "s2-new")
        self.assertEqual(new_row["matched_keywords"], ["allene synthesis", "nickel catalysis"])
        plan = module.build_external_ingest_plan("fixture-review", selected["web_papers"], True)
        actions = {item["paper_key"]: item["action"] for item in plan["items"]}
        self.assertEqual(actions["s2-known"], "use_local")
        self.assertEqual(actions["s2-new"], "download_then_mineru")
        self.assertEqual(plan["downloadable_count"], 1)
        planned_new = next(item for item in plan["items"] if item["paper_key"] == "s2-new")
        self.assertEqual(planned_new["authors"], ["Ada Author"])
        self.assertEqual(planned_new["journal"], "Open Chemistry")
        self.assertIn("discovery abstract", planned_new["abstract"])

        spec = importlib.util.spec_from_file_location("review_external_ingest", EXTERNAL_INGEST)
        self.assertIsNotNone(spec)
        self.assertIsNotNone(spec.loader)
        ingest = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(ingest)
        chosen = ingest.select_items(plan, ["s2-new"], 3)
        self.assertEqual([row["paper_key"] for row in chosen], ["s2-new"])
        explicitly_requested = {
            "items": [
                {
                    "paper_key": f"paper-{index}",
                    "action": "download_then_mineru",
                }
                for index in range(5)
            ]
        }
        chosen = ingest.select_items(
            explicitly_requested,
            [f"paper-{index}" for index in range(5)],
            3,
        )
        self.assertEqual(len(chosen), 5)
        already_imported = {
            "items": [
                {
                    "paper_key": "paper-imported",
                    "action": "use_local",
                    "target_pdf_path": "chem_papers/paper-imported.pdf",
                }
            ]
        }
        self.assertEqual(
            len(ingest.select_items(already_imported, ["paper-imported"], 3)),
            1,
        )
        self.assertEqual(ingest.select_items(already_imported, [], 3), [])
        with self.assertRaises(ValueError):
            ingest.safe_target(self.root, "../outside.pdf")

        page_only = {
            "items": [
                {
                    "paper_key": "10.1000/page-only",
                    "title": "Publisher hosted open article",
                    "year": 2025,
                    "action": "locate_pdf",
                    "open_access_full_text_url": "https://publisher.example/article/full",
                }
            ]
        }
        chosen = ingest.select_items(page_only, [], 3)
        self.assertEqual([row["paper_key"] for row in chosen], ["10.1000/page-only"])
        self.assertEqual(
            ingest.extract_pdf_urls(
                "https://publisher.example/article/full",
                '<html><head><meta name="citation_pdf_url" content="/article/file.pdf"></head>'
                '<body><a href="/article/alternate">Download PDF</a></body></html>',
            ),
            [
                "https://publisher.example/article/file.pdf",
                "https://publisher.example/article/alternate",
            ],
        )
        response = mock.MagicMock()
        response.__enter__.return_value = response
        response.read.return_value = json.dumps(
            {
                "resultList": {
                    "result": [
                        {
                            "doi": "10.1000/page-only",
                            "pmcid": "PMC1234567",
                            "isOpenAccess": "Y",
                            "inEPMC": "Y",
                        }
                    ]
                }
            }
        ).encode("utf-8")
        with mock.patch.object(ingest.urllib.request, "urlopen", return_value=response):
            self.assertEqual(
                ingest.europe_pmc_article_url("10.1000/page-only", 30),
                "https://pmc.ncbi.nlm.nih.gov/articles/PMC1234567/",
            )
        jats_response = mock.MagicMock()
        jats_response.__enter__.return_value = jats_response
        jats_response.read.return_value = (
            b'<article><front><article-meta><article-id pub-id-type="doi">10.1000/page-only</article-id>'
            b'<title-group><article-title>Repository full text article</article-title></title-group>'
            b'<abstract><p>Repository abstract with enough source text.</p></abstract>'
            b'</article-meta></front><body><sec><title>Results</title>'
            b'<p>The repository provides the complete result paragraph.</p></sec></body></article>'
        )
        jats_target = self.root / "chem_papers" / "web-imports" / "repository.pdf"
        with mock.patch.object(ingest.urllib.request, "urlopen", return_value=jats_response):
            xml_path, markdown_path = ingest.download_europe_pmc_jats(
                "https://pmc.ncbi.nlm.nih.gov/articles/PMC1234567/",
                jats_target,
                self.root,
                30,
            )
        self.assertTrue(xml_path.exists())
        self.assertIn("## Results", markdown_path.read_text(encoding="utf-8"))
        compact_target = ingest.default_target_relative(page_only["items"][0])
        self.assertTrue(compact_target.endswith(".pdf"))
        self.assertLessEqual(len(Path(compact_target).stem), 64)
        batch_targets = [
            self.root / "chem_papers" / "web-imports" / "one.pdf",
            self.root / "chem_papers" / "web-imports" / "two.pdf",
        ]
        mineru_command = ingest.build_mineru_command(
            sys.executable,
            MINERU_PARSER,
            self.root,
            batch_targets,
            7,
        )
        self.assertEqual(mineru_command.count("--pdf"), 2)
        self.assertEqual(mineru_command[mineru_command.index("--batch-size") + 1], "7")

        parser_spec = importlib.util.spec_from_file_location("review_mineru_batch", MINERU_PARSER)
        self.assertIsNotNone(parser_spec)
        self.assertIsNotNone(parser_spec.loader)
        mineru = importlib.util.module_from_spec(parser_spec)
        sys.modules[parser_spec.name] = mineru
        try:
            parser_spec.loader.exec_module(mineru)
        finally:
            sys.modules.pop(parser_spec.name, None)
        for target in batch_targets:
            target.parent.mkdir(parents=True, exist_ok=True)
            target.write_bytes(b"%PDF-1.4\n")
        selected_jobs = mineru.discover_selected_jobs(
            [batch_targets[0], batch_targets[1], batch_targets[0]],
            self.root / "chem_papers",
            self.root / "mineru-outputs",
            False,
        )
        self.assertEqual([job.index for job in selected_jobs], [1, 2])
        self.assertEqual(len({job.data_id for job in selected_jobs}), 2)

        metadata_dir = self.root / "review-library" / "metadata" / "papers"
        registry_path = self.root / "review-library" / "registry" / "papers.jsonl"
        write_json(
            metadata_dir / "P900.metadata.json",
            {
                "paper_id": "P900",
                "title": {"value": "A title with C O 2", "confidence": 0.5},
                "authors": {"value": ["A. Author"], "confidence": 0.8},
                "year": {"value": 2025, "confidence": 0.7},
                "journal": {"value": None, "confidence": 0.0},
                "doi": {"value": "10.1000", "confidence": 0.4},
                "abstract": {"value": "Abstract", "confidence": 0.8},
                "structured_tags": {"value": {}, "confidence": 0.0},
                "extraction": {"notes": []},
                "quality": {"missing_fields": [], "warnings": ["missing_doi"]},
            },
        )
        registry_path.parent.mkdir(parents=True, exist_ok=True)
        registry_path.write_text(
            json.dumps({"paper_id": "P900", "title": "Old", "year": 2025, "doi": None}) + "\n",
            encoding="utf-8",
        )
        repaired, fields = ingest.reconcile_bibliographic_metadata(
            self.root,
            "P900",
            json.loads((metadata_dir / "P900.metadata.json").read_text(encoding="utf-8")),
            {
                "title": "A title with CO2",
                "authors": ["Ada Author", "Ben Researcher"],
                "year": 2025,
                "journal": "Open Chemistry",
                "doi": "https://doi.org/10.1000/example",
                "abstract": "An externally supplied abstract with bibliographic provenance.",
            },
        )
        self.assertEqual(fields, ["title", "doi", "authors", "journal", "abstract"])
        self.assertEqual(repaired["doi"]["value"], "10.1000/example")
        self.assertEqual(repaired["authors"]["value"], ["Ada Author", "Ben Researcher"])
        registry = json.loads(registry_path.read_text(encoding="utf-8"))
        self.assertEqual(registry["title"], "A title with CO2")
        self.assertEqual(registry["doi"], "10.1000/example")
        self.assertEqual(registry["journal"], "Open Chemistry")

    def test_matrix_blocks_missing_source_and_draft_blocks_wrong_evidence_owner(self) -> None:
        matrix_path = self.project / "01_matrix_outline" / "literature_matrix.json"
        matrix = json.loads(matrix_path.read_text(encoding="utf-8"))
        matrix["papers"][0]["evidence_anchors"][0]["source_path"] = "missing-source.md"
        write_json(matrix_path, matrix)
        result = run(
            MATRIX_VALIDATOR,
            "--review-root",
            self.root,
            "--project-id",
            self.project_id,
        )
        self.assertNotEqual(result.returncode, 0)
        self.assertIn("source_path does not exist", result.stdout)

        self._write_matrix_and_metadata()
        write_json(
            self.project / "01_matrix_outline" / "section_blueprint.json",
            {"coverage_contract": {"dimensions": []}, "sections": [{"section_id": "sec1"}]},
        )
        write_json(
            self.project / "02_section_drafting" / "section_drafts.json",
            {
                "front_matter": self._front_matter(),
                "sections": [
                    {
                        "section_id": "sec1",
                        "title": "Evidence ownership",
                        "paragraphs": [
                            {
                                "paragraph_id": "sec1-p1",
                                "markdown": "The first paper reports the focal result [@P001].",
                                "evidence_ids": ["P002-E01"],
                            }
                        ],
                    }
                ],
            },
        )
        result = run(
            DRAFT_VALIDATOR,
            "--review-root",
            self.root,
            "--project-id",
            self.project_id,
        )
        self.assertNotEqual(result.returncode, 0)
        self.assertIn("evidence belongs to uncited papers P002", result.stdout)

    def test_matrix_blocks_source_excerpt_that_is_only_a_paraphrase(self) -> None:
        matrix_path = self.project / "01_matrix_outline" / "literature_matrix.json"
        matrix = json.loads(matrix_path.read_text(encoding="utf-8"))
        matrix["papers"][0]["evidence_anchors"][0]["source_excerpt"] = (
            "A broad transformation with useful scope was established by the authors."
        )
        write_json(matrix_path, matrix)
        result = run(
            MATRIX_VALIDATOR,
            "--review-root",
            self.root,
            "--project-id",
            self.project_id,
        )
        self.assertNotEqual(result.returncode, 0)
        self.assertIn("source_excerpt is not found in the recorded source", result.stdout)

    def test_matrix_blocks_paper_title_used_as_full_text_evidence(self) -> None:
        matrix_path = self.project / "01_matrix_outline" / "literature_matrix.json"
        matrix = json.loads(matrix_path.read_text(encoding="utf-8"))
        paper = matrix["papers"][0]
        source_path = self.root / "review-library" / "sources" / "P001.md"
        source_path.write_text(
            f"# {paper['title']}\n\nThe transformation and representative scope are reported.\n",
            encoding="utf-8",
        )
        paper["evidence_anchors"][0]["source_excerpt"] = paper["title"]
        write_json(matrix_path, matrix)
        result = run(
            MATRIX_VALIDATOR,
            "--review-root",
            self.root,
            "--project-id",
            self.project_id,
        )
        self.assertNotEqual(result.returncode, 0)
        self.assertIn("source_excerpt is front matter", result.stdout)

    def test_draft_validator_rejects_invalid_figure_candidates_shape(self) -> None:
        write_json(
            self.project / "01_matrix_outline" / "section_blueprint.json",
            {"coverage_contract": {"dimensions": []}, "sections": [{"section_id": "sec1"}]},
        )
        write_json(
            self.project / "02_section_drafting" / "section_drafts.json",
            {
                "front_matter": self._front_matter(),
                "sections": [
                    {
                        "section_id": "sec1",
                        "title": "Evidence",
                        "paragraphs": [
                            {
                                "paragraph_id": "sec1-p1",
                                "markdown": "The paper reports the focal result [@P001].",
                                "evidence_ids": ["P001-E01"],
                            }
                        ],
                    }
                ],
            },
        )
        write_json(
            self.project / "02_section_drafting" / "figure_candidates.json",
            {"candidates": [], "skip_reason": "not a supported structure"},
        )
        result = run(DRAFT_VALIDATOR, "--review-root", self.root, "--project-id", self.project_id)
        self.assertNotEqual(result.returncode, 0)
        self.assertIn("figure_candidates.json must be a list", result.stdout)

    def test_figure_selector_accepts_object_section_tasks_and_uses_blueprint(self) -> None:
        write_json(
            self.project / "00_discovery" / "selected_discovery_results.json",
            {"local_papers": [{"paper_id": "P001"}]},
        )
        write_json(
            self.project / "01_matrix_outline" / "section_blueprint.json",
            {
                "sections": [
                    {
                        "section_id": "sec1",
                        "title": "Polymer Microstructure and Enzymatic Accessibility",
                        "section_thesis": "Compare how amorphous and crystalline regions affect degradation.",
                        "major_papers": [],
                    }
                ]
            },
        )
        write_json(
            self.project / "02_section_drafting" / "section_tasks.json",
            {
                "project_id": self.project_id,
                "tasks": [
                    {
                        "section_id": "sec1",
                        "title": "Polymer Microstructure and Enzymatic Accessibility",
                        "assigned_papers": ["P001"],
                    }
                ],
            },
        )
        image_path = self.root / "review-library" / "figures" / "P001-scheme.png"
        image_path.parent.mkdir(parents=True, exist_ok=True)
        image_path.write_bytes(b"fixture")
        content_path = image_path.parent / "P001_content_list.json"
        write_json(
            content_path,
            [
                {
                    "type": "image",
                    "img_path": image_path.name,
                    "image_caption": "Figure 1. Polymer microstructure controls enzymatic degradation",
                    "page_idx": 3,
                }
            ],
        )
        metadata_path = self.root / "review-library" / "metadata" / "papers" / "P001.metadata.json"
        metadata = json.loads(metadata_path.read_text(encoding="utf-8"))
        metadata["source_paths"] = {
            "content_list": str(content_path),
            "extracted_dir": str(image_path.parent),
        }
        write_json(metadata_path, metadata)
        write_json(
            self.project / "02_section_drafting" / "paper_figure_inventory.json",
            {
                "papers": [
                    {
                        "paper_id": "P001",
                        "title": "Polymer degradation study",
                        "top_candidates": [
                            {
                                "paper_id": "P001",
                                "source_label": "Figure 1",
                                "source_type": "image",
                                "source_caption_text": "Polymer microstructure controls enzymatic degradation",
                                "source_image_path": str(image_path),
                                "inventory_score": 20,
                            }
                        ],
                    }
                ]
            },
        )
        result = run(
            FIGURE_SELECTOR,
            "--review-root",
            self.root,
            "--project-id",
            self.project_id,
            "--max-total",
            "1",
        )
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        selected = json.loads(
            (self.project / "02_section_drafting" / "figure_candidates.json").read_text(encoding="utf-8")
        )
        inventory = json.loads(
            (self.project / "02_section_drafting" / "paper_figure_inventory.json").read_text(encoding="utf-8")
        )
        paper_candidates = json.loads(
            (self.project / "02_section_drafting" / "paper_figure_candidates.json").read_text(encoding="utf-8")
        )
        self.assertEqual(len(selected), 1)
        self.assertEqual(inventory["candidate_count"], 1)
        self.assertEqual(len(inventory["candidates"]), 1)
        self.assertEqual(len(inventory["papers"][0]["candidates"]), 1)
        self.assertEqual(paper_candidates["inventory_candidate_count"], 1)
        self.assertEqual(len(paper_candidates["candidates"]), 1)
        self.assertEqual(selected[0]["section_id"], "sec1")
        self.assertEqual(selected[0]["source_page_review_status"], "pending")
        self.assertEqual(selected[0]["reuse_rights"]["status"], "pending")

    def test_source_figure_requires_and_records_page_completeness_review(self) -> None:
        image_path = self.root / "review-library" / "figures" / "P001-scheme.png"
        image_path.parent.mkdir(parents=True, exist_ok=True)
        image_path.write_bytes(b"fixture image bytes")
        candidate = {
            "paper_id": "P001",
            "section_id": "sec1",
            "manuscript_selected": True,
            "reader_job": "Show the proposed pathway that the surrounding comparison discusses.",
            "placement_rationale": "Place beside the mechanism discussion where its labels are explained.",
            "reuse_basis": "Source figure reused unchanged for internal research review with attribution.",
            "reuse_rights": {
                "status": "verified",
                "basis": "CC BY 4.0",
                "license_url_or_permission_record": "https://creativecommons.org/licenses/by/4.0/",
                "source_locator": "Article license statement and Scheme 1 credit line on page 4",
                "third_party_material_checked": True,
                "adaptation": "unchanged",
                "attribution_text": "Reproduced from Example et al. under CC BY 4.0.",
            },
            "source_label": "Scheme 1",
            "source_page_hint": "page 4",
            "source_caption_text": "(A) Proposed pathway",
            "source_image_path": str(image_path),
            "source_verification_note": "Checked Scheme 1 on page 4 against the source PDF.",
        }
        figures_path = self.project / "02_section_drafting" / "figure_candidates.json"
        write_json(figures_path, [candidate])
        result = run(
            FIGURE_REDRAW,
            "--review-root",
            self.root,
            "--project-id",
            self.project_id,
            "--use-source",
            "--require-usable",
        )
        self.assertNotEqual(result.returncode, 0)
        self.assertIn("No usable figure", result.stdout + result.stderr)

        candidate["source_page_review_status"] = "passed"
        candidate["source_completeness"] = "complete"
        write_json(figures_path, [candidate])
        result = run(
            FIGURE_REDRAW,
            "--review-root",
            self.root,
            "--project-id",
            self.project_id,
            "--use-source",
            "--require-usable",
        )
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        fidelity_path = self.project / "03_figure_redraw" / "figure_fidelity_review.json"
        fidelity = json.loads(fidelity_path.read_text(encoding="utf-8"))
        self.assertEqual(fidelity["figures"][0]["verdict"], "passed")
        self.assertTrue(
            fidelity["figures"][0]["completeness_signals"]["panel_marker_detected"]
        )

        spec = importlib.util.spec_from_file_location("review_project_status", PROJECT_STATUS)
        self.assertIsNotNone(spec)
        self.assertIsNotNone(spec.loader)
        module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(module)
        figure_stage = next(stage for stage in module.STAGES if stage["id"] == "figure_redraw")
        self.assertTrue(module.stage_status(self.project, figure_stage)["complete"])
        fidelity["figures"][0]["verdict"] = "pending"
        write_json(fidelity_path, fidelity)
        status = module.stage_status(self.project, figure_stage)
        self.assertFalse(status["complete"])
        self.assertIn("figure_fidelity_not_passed:F001", status["semantic_issues"])

        skip_path = self.project / "03_figure_redraw" / "skip_reason.md"
        skip_path.write_text("", encoding="utf-8")
        self.assertFalse(module.stage_status(self.project, figure_stage)["skipped_by_user"])

        result = run(
            PROJECT_STATUS,
            "--review-root",
            self.root,
            "--project-id",
            self.project_id,
            "--require-complete",
        )
        self.assertNotEqual(result.returncode, 0)

    def test_figure_selector_recovers_a_missing_extracted_image_from_the_source_pdf(self) -> None:
        import fitz

        write_json(
            self.project / "00_discovery" / "selected_discovery_results.json",
            {"local_papers": [{"paper_id": "P001"}]},
        )
        write_json(
            self.project / "01_matrix_outline" / "section_blueprint.json",
            {
                "sections": [
                    {
                        "section_id": "sec1",
                        "title": "Mechanistic pathway comparison",
                        "section_thesis": "Compare the proposed catalytic pathways.",
                        "major_papers": [],
                    }
                ]
            },
        )
        source_dir = self.root / "review-library" / "sources"
        source_dir.mkdir(parents=True, exist_ok=True)
        pdf_path = source_dir / "P001.pdf"
        with fitz.open() as document:
            page = document.new_page(width=300, height=200)
            page.draw_rect(fitz.Rect(50, 50, 250, 150), color=(0, 0, 0))
            page.insert_text((75, 100), "Catalytic pathway", fontsize=16)
            document.save(pdf_path)
        content_path = source_dir / "P001_content_list.json"
        write_json(
            content_path,
            [
                {
                    "type": "image",
                    "img_path": "missing-extraction.png",
                    "image_caption": "Figure 2. Proposed catalytic pathway",
                    "page_idx": 0,
                    "bbox": [50, 50, 250, 150],
                }
            ],
        )
        markdown_path = source_dir / "P001.md"
        markdown_path.write_text(
            "Licensed under Creative Commons CC BY 4.0.\n", encoding="utf-8"
        )
        metadata_path = self.root / "review-library" / "metadata" / "papers" / "P001.metadata.json"
        metadata = json.loads(metadata_path.read_text(encoding="utf-8"))
        metadata["source_paths"] = {
            "pdf": str(pdf_path),
            "markdown": str(markdown_path),
            "content_list": str(content_path),
            "extracted_dir": str(source_dir / "missing-extracted-dir"),
        }
        write_json(metadata_path, metadata)

        result = run(
            FIGURE_SELECTOR,
            "--review-root",
            self.root,
            "--project-id",
            self.project_id,
            "--max-total",
            "1",
        )
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        selected = json.loads(
            (self.project / "02_section_drafting" / "figure_candidates.json").read_text(
                encoding="utf-8"
            )
        )
        self.assertEqual(len(selected), 1)
        crop = Path(selected[0]["source_image_path"])
        self.assertTrue(crop.is_file())
        self.assertEqual(crop.read_bytes()[:8], b"\x89PNG\r\n\x1a\n")
        self.assertEqual(selected[0]["source_resolution_status"], "materialized_pdf_crop")
        self.assertEqual(selected[0]["source_page_review_status"], "pending")

    def test_source_figure_floor_counts_only_unique_unchanged_figures_from_cited_papers(self) -> None:
        spec = importlib.util.spec_from_file_location("review_final_audit_source_figures", AUDIT)
        module = importlib.util.module_from_spec(spec)
        assert spec.loader is not None
        spec.loader.exec_module(module)

        inserted = []
        figures = []
        image_paths = []
        for index in range(1, 4):
            figure_id = f"F{index:03d}"
            inserted_path = f"figures/source-{index}.png"
            image_paths.append(inserted_path)
            inserted.append(
                {
                    "figure_id": figure_id,
                    "mode": "source_verified",
                    "inserted_path": inserted_path,
                }
            )
            figures.append(
                {
                    "figure_id": figure_id,
                    "paper_id": "P001",
                    "source_label": f"Figure {index}",
                    "status": "source_verified",
                    "verification_status": "passed",
                    "reuse_rights": {
                        "status": "verified",
                        "basis": "CC BY 4.0",
                        "license_url_or_permission_record": "https://creativecommons.org/licenses/by/4.0/",
                        "source_locator": f"Figure {index} credit line",
                        "third_party_material_checked": True,
                        "adaptation": "unchanged",
                        "attribution_text": "Reproduced from the cited paper under CC BY 4.0.",
                    },
                }
            )
        write_json(
            self.project / "04_first_draft" / "figure_insertion_report.json",
            {"inserted": inserted},
        )
        manifest_path = self.project / "03_figure_redraw" / "redrawn_figure_manifest.json"
        write_json(manifest_path, {"figures": figures})

        count, issues = module.verified_source_figure_usage(
            self.project, image_paths, {"P001"}
        )
        self.assertEqual((count, issues), (3, []))

        figures[2]["reuse_rights"]["adaptation"] = "adapted"
        write_json(manifest_path, {"figures": figures})
        count, issues = module.verified_source_figure_usage(
            self.project, image_paths, {"P001"}
        )
        self.assertEqual(count, 2)
        self.assertIn("inserted_source_figure_provenance_incomplete:F003", issues)

        figures[2]["reuse_rights"]["adaptation"] = "unchanged"
        inserted.append(dict(inserted[0]))
        write_json(manifest_path, {"figures": figures})
        write_json(
            self.project / "04_first_draft" / "figure_insertion_report.json",
            {"inserted": inserted},
        )
        count, issues = module.verified_source_figure_usage(
            self.project, image_paths, {"P001"}
        )
        self.assertEqual(count, 3)
        self.assertIn("duplicate_inserted_source_figure:F001", issues)

    def test_comprehensive_source_figure_floor_rejects_a_bound_table(self) -> None:
        spec = importlib.util.spec_from_file_location("review_final_audit_non_table", AUDIT)
        module = importlib.util.module_from_spec(spec)
        assert spec.loader is not None
        spec.loader.exec_module(module)

        write_json(
            self.project / "00_discovery" / "topic_contract.json",
            {"review_profile": "comprehensive"},
        )
        source_pdf = self.root / "review-library" / "pdf" / "P001.pdf"
        source_pdf.parent.mkdir(parents=True, exist_ok=True)
        source_pdf.write_bytes(b"source pdf fixture")
        inserted_asset = self.project / "05_final_audit" / "figures" / "source-table.png"
        inserted_asset.parent.mkdir(parents=True, exist_ok=True)
        inserted_asset.write_bytes(b"source table image")
        candidate = {
            "inventory_candidate_id": "P001-table-1",
            "paper_id": "P001",
            "source_label": "Table 1",
            "source_type": "table",
            "source_caption_text": "Conditions and yields",
            "source_pdf_sha256": module.file_sha256(source_pdf),
            "source_page_index": 2,
            "source_bbox": [10, 20, 300, 400],
        }
        write_json(
            self.project / "02_section_drafting" / "paper_figure_inventory.json",
            {"candidates": [candidate]},
        )
        source = {
            **candidate,
            "figure_id": "F001",
            "status": "source_verified",
            "verification_status": "passed",
            "source_pdf": str(source_pdf),
            "source_image": str(inserted_asset),
            "accepted_image_sha256": module.file_sha256(inserted_asset),
            "source_completeness": "complete",
            "source_page_review_status": "passed",
            "reader_job": "Compare reported conditions.",
            "placement_rationale": "Placed beside the conditions discussion.",
            "manuscript_callout": "Table 1 provides the source comparison.",
            "reuse_rights": {
                "status": "verified",
                "basis": "CC BY 4.0",
                "license_url_or_permission_record": "https://creativecommons.org/licenses/by/4.0/",
                "source_locator": "Table 1 credit line",
                "third_party_material_checked": True,
                "adaptation": "unchanged",
                "attribution_text": "Reproduced under CC BY 4.0.",
            },
        }
        write_json(
            self.project / "03_figure_redraw" / "redrawn_figure_manifest.json",
            {"figures": [source]},
        )
        write_json(
            self.project / "05_final_audit" / "figure_insertion_report.json",
            {
                "inserted": [
                    {
                        "figure_id": "F001",
                        "mode": "source_verified",
                        "inserted_path": "figures/source-table.png",
                    }
                ]
            },
        )
        count, issues = module.verified_source_figure_usage(
            self.project,
            ["figures/source-table.png"],
            {"P001"},
            "Table 1 provides the source comparison.",
        )
        self.assertEqual(count, 0)
        self.assertTrue(any("non-table image/chart" in issue for issue in issues))

    def test_docx_status_requires_a_pdf_rendered_from_docx_and_real_page_images(self) -> None:
        final_dir = self.project / "05_final_audit"
        final_dir.mkdir(parents=True, exist_ok=True)
        (final_dir / "final_draft.md").write_text("# Review\n", encoding="utf-8")
        (final_dir / "final_draft.docx").write_bytes(b"docx-fixture")
        (final_dir / "final_draft.pdf").write_bytes(b"pdf-fixture")
        write_json(final_dir / "format_scan.json", {"blocking_issues": []})
        write_json(final_dir / "docx_audit.json", {"blocking_issues": [], "render_qa": "passed"})
        write_json(
            final_dir / "render_qa_report.json",
            {
                "render_status": "passed",
                "inspection_status": "passed",
                "renderer": "hand-written-report",
                "output_pdf": "final_draft.pdf",
                "page_count": 1,
            },
        )
        spec = importlib.util.spec_from_file_location("review_project_status_render", PROJECT_STATUS)
        module = importlib.util.module_from_spec(spec)
        assert spec.loader is not None
        spec.loader.exec_module(module)
        stage = next(row for row in module.STAGES if row["id"] == "docx_export")
        status = module.stage_status(self.project, stage)
        self.assertFalse(status["complete"])
        self.assertIn("final_docx_not_canonical_render_input", status["semantic_issues"])
        self.assertIn("rendered_page_image_count_mismatch", status["semantic_issues"])

        page = final_dir / "rendered_pages" / "page-1.png"
        page.parent.mkdir(parents=True, exist_ok=True)
        page.write_bytes(b"png-fixture")
        docx_hash = hashlib.sha256((final_dir / "final_draft.docx").read_bytes()).hexdigest()
        pdf_hash = hashlib.sha256((final_dir / "final_draft.pdf").read_bytes()).hexdigest()
        page_hash = hashlib.sha256(page.read_bytes()).hexdigest()
        write_json(
            final_dir / "render_qa_report.json",
            {
                "input_docx": "final_draft.docx",
                "output_pdf": "final_draft.pdf",
                "pages_dir": "rendered_pages",
                "render_status": "passed",
                "inspection_status": "passed",
                "renderer": "fixture-renderer",
                "input_docx_sha256": docx_hash,
                "output_pdf_sha256": pdf_hash,
                "page_count": 1,
                "page_images": ["rendered_pages/page-1.png"],
                "page_artifacts": [
                    {
                        "page_number": 1,
                        "path": "rendered_pages/page-1.png",
                        "sha256": page_hash,
                    }
                ],
                "page_inspections": [
                    {
                        "page_number": 1,
                        "page_sha256": page_hash,
                        "verdict": "passed",
                        "observation": "The single rendered page has a visible title, intact margins, and no clipped content.",
                    }
                ],
            },
        )
        status = module.stage_status(self.project, stage)
        self.assertTrue(status["complete"], status["semantic_issues"])

    def test_run_recorder_verifies_declared_artifacts_but_status_does_not_gate_on_history(self) -> None:
        artifact = self.project / "00_discovery" / "receipt-probe.txt"
        result = run(
            RUN_AND_RECORD,
            "--review-root",
            self.root,
            "--project-id",
            self.project_id,
            "--stage",
            "discovery",
            "--artifact",
            "00_discovery/receipt-probe.txt",
            "--",
            sys.executable,
            "-c",
            f"from pathlib import Path; Path({str(artifact)!r}).write_text('recorded', encoding='utf-8')",
        )
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        event = json.loads(
            (self.project / "run_events.jsonl").read_text(encoding="utf-8").splitlines()[-1]
        )
        self.assertEqual(event["event_version"], 2)
        self.assertTrue(event["artifact_receipts"][0]["after"]["exists"])
        self.assertTrue(event["artifact_receipts"][0]["changed_by_command"])
        self.assertTrue(event["artifact_receipts"][0]["content_changed_by_command"])

        missing = run(
            RUN_AND_RECORD,
            "--review-root",
            self.root,
            "--project-id",
            self.project_id,
            "--stage",
            "discovery",
            "--artifact",
            "00_discovery/not-created.txt",
            "--",
            sys.executable,
            "-c",
            "print('no artifact created')",
        )
        self.assertEqual(missing.returncode, 2)
        self.assertIn("declared artifact missing", missing.stderr)

        spec = importlib.util.spec_from_file_location("review_project_status_truth", PROJECT_STATUS)
        self.assertIsNotNone(spec)
        self.assertIsNotNone(spec.loader)
        module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(module)
        placeholder_events = [
            {
                "event_version": 2,
                "project_id": self.project_id,
                "stage": stage,
                "started_at": "2026-07-21T00:00:00Z",
                "finished_at": "2026-07-21T00:00:01Z",
                "cwd": str(self.root),
                "command": [sys.executable, "-c", "print('passed')"],
                "exit_code": 0,
                "artifact_receipts": [
                    {"artifact": "probe", "after": {"exists": True}}
                ],
            }
            for stage in ("final_audit", "docx_export")
        ]
        (self.project / "run_events.jsonl").write_text(
            "\n".join(json.dumps(row) for row in placeholder_events) + "\n",
            encoding="utf-8",
        )
        (self.project / "run_record.md").write_text(
            "Generated from `run_events.jsonl`\n", encoding="utf-8"
        )
        issues = module.run_record_issues(
            self.project,
            [
                {"id": "final_audit", "complete": True},
                {"id": "docx_export", "complete": True},
            ],
        )
        self.assertEqual(issues, [])

    def test_reference_metadata_requires_doi_or_journal_locator(self) -> None:
        spec = importlib.util.spec_from_file_location("review_final_audit", AUDIT)
        self.assertIsNotNone(spec)
        self.assertIsNotNone(spec.loader)
        module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(module)
        rows = {"P001": {"paper_id": "P001"}}
        self.assertEqual(module.actual_incomplete_references(self.root, ["P001"], rows), [])

        metadata_path = self.root / "review-library" / "metadata" / "papers" / "P001.metadata.json"
        metadata = json.loads(metadata_path.read_text(encoding="utf-8"))
        metadata["doi"] = {"value": ""}
        write_json(metadata_path, metadata)
        issues = module.actual_incomplete_references(self.root, ["P001"], rows)
        self.assertEqual(issues[0]["missing"], ["doi_or_volume_and_page_locator"])

        metadata["volume"] = {"value": "12"}
        metadata["article_number"] = {"value": "104321"}
        write_json(metadata_path, metadata)
        self.assertEqual(module.actual_incomplete_references(self.root, ["P001"], rows), [])

    def test_final_audit_blocks_phantom_table_and_mojibake(self) -> None:
        spec = importlib.util.spec_from_file_location("review_final_audit_integrity", AUDIT)
        self.assertIsNotNone(spec)
        self.assertIsNotNone(spec.loader)
        module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(module)

        draft_path = self.project / "04_first_draft" / "first_draft.md"
        draft_path.parent.mkdir(parents=True, exist_ok=True)
        base = (
            "# Review\n\n## Abstract\n\nA bounded abstract.\n\n"
            "## Keywords\n\npolymer; recycling; evidence\n\n"
            "## 1. Introduction\n\nTable 1 summarizes the comparison [1]. PET\ufffd values are corrupted.\n\n"
            "## 2. Conclusion and Outlook\n\nThe comparison remains bounded [1].\n\n"
            "## References\n\n[1] Example reference.\n"
        )
        draft_path.write_text(base, encoding="utf-8")
        report = module.scan_draft(self.project, "preflight")
        self.assertIn("manuscript_references_missing_table", report["blocking_issues"])
        self.assertIn("mojibake_or_replacement_characters_present", report["blocking_issues"])

        with_table = base.replace(
            "Table 1 summarizes the comparison [1].",
            "Table 1 summarizes the comparison [1].\n\n| Method | Result |\n|---|---|\n| A | B |",
        ).replace("PET\ufffd", "PET")
        draft_path.write_text(with_table, encoding="utf-8")
        report = module.scan_draft(self.project, "preflight")
        self.assertNotIn("manuscript_references_missing_table", report["blocking_issues"])
        self.assertNotIn("mojibake_or_replacement_characters_present", report["blocking_issues"])

    def test_comprehensive_release_has_one_coarse_product_floor(self) -> None:
        write_json(
            self.project / "00_discovery" / "topic_contract.json",
            {"review_profile": "comprehensive", "topic": "A broad review"},
        )
        final = self.project / "05_final_audit" / "final_draft.md"
        final.parent.mkdir(parents=True, exist_ok=True)
        final.write_text(
            "# Review\n\n## Abstract\n\nA bounded abstract.\n\n"
            "**Keywords:** evidence; review; methods\n\n"
            "## 1. Introduction\n\nA short supported statement [1].\n\n"
            "## 2. Conclusion and Outlook\n\nA bounded conclusion [1].\n\n"
            "## References\n\n[1] Example reference.\n",
            encoding="utf-8",
        )
        spec = importlib.util.spec_from_file_location("review_final_audit_floor", AUDIT)
        module = importlib.util.module_from_spec(spec)
        assert spec.loader is not None
        spec.loader.exec_module(module)
        report = module.scan_draft(self.project, "release")
        floor_issues = report["delivery_floor_issues"]
        self.assertEqual(report["review_profile"], "comprehensive")
        self.assertEqual(len(floor_issues), 4)
        self.assertTrue(all(issue in report["blocking_issues"] for issue in floor_issues))
        self.assertTrue(any(":word_like_count:" in issue for issue in floor_issues))
        self.assertTrue(any(":reference_count:" in issue for issue in floor_issues))
        self.assertTrue(any(":table_count:" in issue for issue in floor_issues))
        self.assertFalse(any(":figure_count:" in issue for issue in floor_issues))
        self.assertTrue(any(":source_figure_count:" in issue for issue in floor_issues))

        status_spec = importlib.util.spec_from_file_location(
            "review_project_status_floor", PROJECT_STATUS
        )
        status_module = importlib.util.module_from_spec(status_spec)
        assert status_spec.loader is not None
        status_spec.loader.exec_module(status_module)
        stage = next(row for row in status_module.STAGES if row["id"] == "final_audit")
        status = status_module.stage_status(self.project, stage)
        self.assertTrue(
            all(
                any(
                    candidate.startswith(issue.rsplit(":", 1)[0] + ":")
                    for candidate in status["semantic_issues"]
                )
                for issue in floor_issues
            ),
            status["semantic_issues"],
        )

    def test_disguised_backmatter_cannot_inflate_body_or_citation_counts(self) -> None:
        write_json(
            self.project / "00_discovery" / "topic_contract.json",
            {"review_profile": "comprehensive", "topic": "A broad review"},
        )
        padding = " ".join(["padding"] * 9000)
        fake_references = "\n".join(
            f"[{index}] Fabricated-looking reference entry." for index in range(1, 31)
        )
        final = self.project / "05_final_audit" / "final_draft.md"
        final.parent.mkdir(parents=True, exist_ok=True)
        final.write_text(
            "# Review\n\n## Abstract\n\nA bounded abstract.\n\n"
            "**Keywords:** evidence; review; methods\n\n"
            "## 1. Introduction\n\nA short supported statement [1].\n\n"
            "## 2. Conclusion and Outlook\n\nA bounded conclusion [1].\n\n"
            f"**References**\n\n{padding}\n\n{fake_references}\n\n"
            f"## References\n\n{fake_references}\n",
            encoding="utf-8",
        )
        spec = importlib.util.spec_from_file_location("review_final_audit_backmatter", AUDIT)
        module = importlib.util.module_from_spec(spec)
        assert spec.loader is not None
        spec.loader.exec_module(module)
        report = module.scan_draft(self.project, "release")
        self.assertLess(report["delivery_metrics"]["word_like_count"], 100)
        self.assertEqual(report["delivery_metrics"]["reference_count"], 1)
        self.assertIn("noncanonical_or_duplicated_backmatter", report["blocking_issues"])

    def test_final_audit_catches_inventory_count_mismatch_without_forcing_full_disposition(self) -> None:
        spec = importlib.util.spec_from_file_location("review_final_audit_figures", AUDIT)
        self.assertIsNotNone(spec)
        self.assertIsNotNone(spec.loader)
        module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(module)
        write_json(
            self.project / "02_section_drafting" / "paper_figure_inventory.json",
            {
                "candidate_count": 2,
                "papers": [
                    {"paper_id": "P001", "candidate_count": 2, "top_candidates": [{}, {}]}
                ],
            },
        )
        write_json(
            self.project / "02_section_drafting" / "paper_figure_candidates.json",
            {"total_inventory_figures": 0, "candidates": []},
        )
        issues = module.figure_inventory_consistency_issues(self.project)
        self.assertIn("paper_figure_candidate_count_mismatch", issues)
        self.assertNotIn("source_figure_inventory_not_reviewed", issues)

    def test_reference_doi_must_match_linked_source_pdf(self) -> None:
        spec = importlib.util.spec_from_file_location("review_final_audit", AUDIT)
        self.assertIsNotNone(spec)
        self.assertIsNotNone(spec.loader)
        module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(module)

        pdf_path = self.root / "review-library" / "sources" / "P001.pdf"
        pdf_path.parent.mkdir(parents=True, exist_ok=True)
        writer = PdfWriter()
        writer.add_blank_page(width=72, height=72)
        writer.add_metadata({"/Subject": "Article DOI 10.1000/source-doi"})
        with pdf_path.open("wb") as stream:
            writer.write(stream)

        metadata_path = self.root / "review-library" / "metadata" / "papers" / "P001.metadata.json"
        metadata = json.loads(metadata_path.read_text(encoding="utf-8"))
        metadata["doi"] = {"value": "10.1000/wrong-doi"}
        metadata["source_paths"] = {"pdf": str(pdf_path)}
        write_json(metadata_path, metadata)
        rows = {"P001": {"paper_id": "P001"}}
        conflicts = module.reference_metadata_source_conflicts(self.root, ["P001"], rows)
        self.assertEqual(conflicts[0]["source_dois"], ["10.1000/source-doi"])

        metadata["doi"] = {"value": "https://doi.org/10.1000/source-doi"}
        write_json(metadata_path, metadata)
        self.assertEqual(module.reference_metadata_source_conflicts(self.root, ["P001"], rows), [])
        manuscript = (
            "# Review\n\n## References\n\n"
            "1. Ada Chemist. Example paper. https://doi.org/10.1000/stale-doi\n"
        )
        delivered = module.manuscript_reference_metadata_conflicts(
            manuscript, self.root, {1: "P001"}, rows
        )
        self.assertEqual(delivered[0]["expected_doi"], "10.1000/source-doi")

    def test_stable_citations_audit_and_docx_chemistry(self) -> None:
        payload = {
            "front_matter": self._front_matter(),
            "sections": [
                {
                    "section_id": "sec1",
                    "title": "1. Introduction",
                    "paragraphs": [
                        {
                            "paragraph_id": "sec1-p1",
                            "paragraph_type": "comparison",
                            "markdown": "Palladium carbonylation and nickel electrocarboxylation provide distinct entry points to allenes [@P001; @P002].",
                            "cited_paper_ids": ["P001", "P002"],
                            "evidence_ids": ["P001-E01", "P002-E01"],
                        },
                        {
                            "paragraph_id": "sec1-p2",
                            "paragraph_type": "mechanism",
                            "markdown": "The proposed pathways should remain qualified because their evidentiary bases differ [@P001; @P002]. Table 1 compares this method-level distinction explicitly. CO_2_ incorporation, sp^2^ rehybridization, η1-allenyl binding, and η3-propargyl binding are discussed without converting locants such as C1.\n\n| Method | Distinction |\n|---|---|\n| Palladium | Carbonylation |\n| Nickel | Electrocarboxylation |",
                            "cited_paper_ids": ["P001", "P002"],
                            "evidence_ids": ["P001-E02", "P002-E02"],
                        },
                    ],
                },
                {
                    "section_id": "sec2",
                    "title": "2. Conclusion and Outlook",
                    "paragraphs": [
                        {
                            "paragraph_id": "sec2-p1",
                            "paragraph_type": "limitation_or_gap",
                            "markdown": "Both approaches retain limitations in scope and direct mechanistic verification today [@P001; @P002]. H_2_O and S_N2_ notation exercise explicit subscripts.",
                            "cited_paper_ids": ["P001", "P002"],
                            "evidence_ids": ["P001-E01", "P001-E02", "P002-E01", "P002-E02"],
                        }
                    ],
                },
            ],
        }
        write_json(self.project / "02_section_drafting" / "section_drafts.json", payload)
        p001_metadata_path = (
            self.root / "review-library" / "metadata" / "papers" / "P001.metadata.json"
        )
        p001_metadata = json.loads(p001_metadata_path.read_text(encoding="utf-8"))
        p001_metadata["journal"] = {"value": "46-Journal of Allene Chemistry 2020"}
        write_json(p001_metadata_path, p001_metadata)
        write_json(
            self.project / "01_matrix_outline" / "section_blueprint.json",
            {
                "coverage_contract": {
                    "minimum_manuscript_words": 60,
                    "dimensions": [
                        {
                            "name": "substrate_class",
                            "items": [
                                {
                                    "name": "propargylic carbonates",
                                    "required": True,
                                    "covered_by": ["P001", "P002"],
                                }
                            ],
                        }
                    ],
                },
                "sections": [
                    {
                        "section_id": "sec1",
                        "title": "1. Introduction",
                        "target_words": 41,
                        "paragraph_types_required": ["comparison", "mechanism"],
                        "review_claims": [
                            {"claim_type": "comparison", "supporting_papers": ["P001", "P002"]}
                        ],
                    },
                    {
                        "section_id": "sec2",
                        "title": "2. Conclusion and Outlook",
                        "target_words": 19,
                        "paragraph_types_required": ["limitation_or_gap"],
                        "review_claims": [
                            {"claim_type": "limitation", "supporting_papers": ["P001", "P002"]}
                        ],
                    },
                ],
            },
        )
        for validator in (MATRIX_VALIDATOR, BLUEPRINT_VALIDATOR, DRAFT_VALIDATOR):
            result = run(validator, "--review-root", self.root, "--project-id", self.project_id)
            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        result = run(MERGE, "--review-root", self.root, "--project-id", self.project_id)
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        first = self.project / "04_first_draft" / "first_draft.md"
        text = first.read_text(encoding="utf-8")
        self.assertIn("[1,2]", text)
        self.assertNotIn("[@P", text)
        self.assertNotIn("[P001]", text)
        self.assertIn("*Journal of Allene Chemistry* 2020.", text)
        self.assertNotIn("46-Journal", text)
        self.assertNotIn("2020* 2020", text)

        skip = self.project / "03_figure_redraw" / "skip_reason.md"
        skip.parent.mkdir(parents=True, exist_ok=True)
        skip.write_text("Text-first regression fixture.\n", encoding="utf-8")
        final_dir = self.project / "05_final_audit"
        final_dir.mkdir(parents=True, exist_ok=True)
        final = final_dir / "final_draft.md"
        final.write_text(text, encoding="utf-8")
        result = run(AUDIT, "--review-root", self.root, "--project-id", self.project_id, "--phase", "preflight")
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        queue = json.loads((final_dir / "semantic_audit_queue.json").read_text(encoding="utf-8"))
        self.assertEqual(queue["coverage_mode"], "risk_stratified_with_section_sample")
        self.assertEqual(
            set(queue["required_queue_ids"]),
            {"sec1-p1-a1", "sec1-p2-a1", "sec2-p1-a1"},
        )
        write_json(
            final_dir / "semantic_audit.json",
            {
                "model": "fixture-model",
                "manuscript_sha256": hashlib.sha256(final.read_bytes()).hexdigest(),
                "checks": [
                    {
                        "check_id": "A001",
                        "queue_id": "sec1-p1-a1",
                        "section_id": "sec1",
                        "paragraph_id": "sec1-p1",
                        "text_span": "Palladium carbonylation and nickel electrocarboxylation provide distinct entry points to allenes [@P001; @P002].",
                        "cited_paper_ids": ["P001", "P002"],
                        "evidence_ids": ["P001-E01", "P002-E01"],
                        "source_receipts": self._source_receipts("P001-E01", "P002-E01"),
                        "source_checked": True,
                        "support_scope": "full",
                        "verdict": "supported",
                        "comment": "Both checked excerpts report the transformation, representative substrate scope, quantified yields, and operating conditions; this supports the comparison.",
                    },
                    {
                        "check_id": "A002",
                        "queue_id": "sec1-p2-a1",
                        "section_id": "sec1",
                        "paragraph_id": "sec1-p2",
                        "text_span": "The proposed pathways should remain qualified because their evidentiary bases differ [@P001; @P002]. Table 1 compares this method-level distinction explicitly. CO_2_ incorporation, sp^2^ rehybridization, η1-allenyl binding, and η3-propargyl binding are discussed without converting locants such as C1.",
                        "cited_paper_ids": ["P001", "P002"],
                        "evidence_ids": ["P001-E02", "P002-E02"],
                        "source_receipts": self._source_receipts("P001-E02", "P002-E02"),
                        "source_checked": True,
                        "support_scope": "full",
                        "verdict": "supported",
                        "comment": "Both checked excerpts present the pathways as author-proposed mechanisms based on intermediate trapping evidence, which supports qualified wording.",
                    },
                    {
                        "check_id": "A003",
                        "queue_id": "sec2-p1-a1",
                        "section_id": "sec2",
                        "paragraph_id": "sec2-p1",
                        "text_span": "Both approaches retain limitations in scope and direct mechanistic verification today [@P001; @P002]. H_2_O and S_N2_ notation exercise explicit subscripts.",
                        "cited_paper_ids": ["P001", "P002"],
                        "evidence_ids": ["P001-E02", "P002-E02"],
                        "source_receipts": self._source_receipts("P001-E02", "P002-E02"),
                        "source_checked": True,
                        "support_scope": "full",
                        "verdict": "supported",
                        "comment": "The checked excerpts label the pathways as mechanistic proposals based on intermediate trapping evidence, preserving the stated verification boundary.",
                    },
                ],
                "section_checks": [
                    {"section_id": "sec1", "coverage_complete": True, "comment": "checked"},
                    {"section_id": "sec2", "coverage_complete": True, "comment": "checked"},
                ],
                "changes_made": [],
                "unresolved_blockers": [],
            },
        )
        result = run(AUDIT, "--review-root", self.root, "--project-id", self.project_id, "--phase", "release")
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)

        receipt_checked_payload = json.loads(
            (final_dir / "semantic_audit.json").read_text(encoding="utf-8")
        )
        without_receipts = json.loads(json.dumps(receipt_checked_payload))
        without_receipts["checks"][0].pop("source_receipts")
        write_json(final_dir / "semantic_audit.json", without_receipts)
        result = run(AUDIT, "--review-root", self.root, "--project-id", self.project_id, "--phase", "release")
        self.assertNotEqual(result.returncode, 0)
        self.assertIn(
            "source receipt verification failed for 1 supported checks",
            result.stdout,
        )
        write_json(final_dir / "semantic_audit.json", receipt_checked_payload)

        docx = final_dir / "final_draft.docx"
        result = run(MD2DOCX, "--input", final, "--output", docx)
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        result = run(DOCX_AUDIT, "--input", docx, "--markdown", final)
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        document = Document(docx)
        self.assertEqual(len(document.tables), 1)
        self.assertGreater(sum(bool(run.font.subscript) for p in document.paragraphs for run in p.runs), 0)
        self.assertGreater(sum(bool(run.font.superscript) for p in document.paragraphs for run in p.runs), 0)
        c1_runs = [run for p in document.paragraphs for run in p.runs if "C1" in run.text]
        self.assertTrue(c1_runs)
        self.assertTrue(all(not run.font.subscript and not run.font.superscript for run in c1_runs))
        eta_paragraphs = [p for p in document.paragraphs if "η1" in p.text and "η3" in p.text]
        self.assertTrue(eta_paragraphs)
        eta_superscripts = [
            run.text
            for paragraph in eta_paragraphs
            for run in paragraph.runs
            if run.font.superscript
        ]
        self.assertIn("1", eta_superscripts)
        self.assertIn("3", eta_superscripts)

        write_json(
            final_dir / "semantic_audit.json",
            {
                "model": "fixture-model",
                "manuscript_sha256": hashlib.sha256(final.read_bytes()).hexdigest(),
                "checks": [
                    {
                        "check_id": "A001",
                        "queue_id": "sec1-p1-a1",
                        "section_id": "sec1",
                        "paragraph_id": "sec1-p1",
                        "text_span": "Representative claim from Introduction section reviewed against evidence",
                        "cited_paper_ids": ["P001", "P002"],
                        "evidence_ids": ["P001-E01", "P002-E01"],
                        "verdict": "supported",
                        "comment": "Fixture placeholder must fail.",
                    },
                    {
                        "check_id": "A002",
                        "queue_id": "sec2-p1-a1",
                        "section_id": "sec2",
                        "paragraph_id": "sec2-p1",
                        "text_span": "Both approaches retain limitations in scope and direct mechanistic verification today [@P001; @P002].",
                        "cited_paper_ids": ["P001", "P002"],
                        "evidence_ids": ["P001-E01", "P002-E01"],
                        "verdict": "supported",
                        "comment": "The source describes the stated limitation and its mechanistic boundary.",
                    },
                ],
                "section_checks": [
                    {"section_id": "sec1", "coverage_complete": True, "comment": "checked"},
                    {"section_id": "sec2", "coverage_complete": True, "comment": "checked"},
                ],
                "changes_made": [],
                "unresolved_blockers": [],
            },
        )
        result = run(AUDIT, "--review-root", self.root, "--project-id", self.project_id, "--phase", "release")
        self.assertNotEqual(result.returncode, 0)
        self.assertIn("placeholder text_span", result.stdout)

        semantic_payload = json.loads(
            (final_dir / "semantic_audit.json").read_text(encoding="utf-8")
        )
        semantic_payload["checks"][0]["text_span"] = (
            "A plausible but fabricated mechanistic statement that is absent from the linked paragraph."
        )
        write_json(final_dir / "semantic_audit.json", semantic_payload)
        result = run(AUDIT, "--review-root", self.root, "--project-id", self.project_id, "--phase", "release")
        self.assertNotEqual(result.returncode, 0)
        self.assertIn("text_span is not found in paragraph", result.stdout)

        write_json(
            final_dir / "semantic_audit.json",
            {
                "model": "fixture-model",
                "manuscript_sha256": hashlib.sha256(final.read_bytes()).hexdigest(),
                "checks": [
                    {
                        "check_id": "A001",
                        "queue_id": "sec1-p1-a1",
                        "section_id": "sec1",
                        "paragraph_id": "sec1-p1",
                        "text_span": "Palladium carbonylation and nickel electrocarboxylation provide distinct entry points to allenes [@P001; @P002].",
                        "cited_paper_ids": ["P001", "P002"],
                        "evidence_ids": ["P001-E01", "P002-E01"],
                        "verdict": "supported",
                        "comment": "The source reports both transformations and supports the stated comparison.",
                    },
                    {
                        "check_id": "A002",
                        "queue_id": "sec2-p1-a1",
                        "section_id": "sec2",
                        "paragraph_id": "sec2-p1",
                        "text_span": "Both approaches retain limitations in scope and direct mechanistic verification today [@P001; @P002].",
                        "cited_paper_ids": ["P001", "P002"],
                        "evidence_ids": ["P001-E01", "P002-E01"],
                        "verdict": "supported",
                        "comment": "The source describes the stated limitation and its mechanistic boundary.",
                    },
                ],
                "section_checks": [
                    {"section_id": "sec1", "coverage_complete": True, "comment": "checked"},
                    {"section_id": "sec2", "coverage_complete": True, "comment": "checked"},
                ],
                "changes_made": [],
                "unresolved_blockers": [],
            },
        )
        metadata_path = (
            self.root / "review-library" / "metadata" / "papers" / "P002.metadata.json"
        )
        metadata = json.loads(metadata_path.read_text(encoding="utf-8"))
        metadata["journal"] = {"value": ""}
        write_json(metadata_path, metadata)
        citations_path = self.project / "04_first_draft" / "citations.json"
        citations = json.loads(citations_path.read_text(encoding="utf-8"))
        citations["incomplete_reference_metadata"] = []
        write_json(citations_path, citations)
        result = run(AUDIT, "--review-root", self.root, "--project-id", self.project_id, "--phase", "release")
        self.assertNotEqual(result.returncode, 0)
        self.assertIn("incomplete_reference_metadata", result.stdout)

        metadata["journal"] = {"value": "Journal of Allene Chemistry"}
        write_json(metadata_path, metadata)
        (self.project / "00_discovery" / "screening_validation.json").unlink()
        result = run(AUDIT, "--review-root", self.root, "--project-id", self.project_id, "--phase", "release")
        self.assertNotEqual(result.returncode, 0)
        self.assertIn("missing_upstream_validation:screening", result.stdout)

    def test_free_form_paragraph_mix_and_word_plan_shortfall_are_advisory(self) -> None:
        write_json(
            self.project / "01_matrix_outline" / "section_blueprint.json",
            {
                "coverage_contract": {"suggested_manuscript_words": 5000, "dimensions": []},
                "sections": [
                    {
                        "section_id": "sec1",
                        "title": "A focused narrative",
                        "target_words": 2000,
                        "paragraph_types_required": ["comparison", "limitation_or_gap"],
                    }
                ],
            },
        )
        write_json(
            self.project / "02_section_drafting" / "section_drafts.json",
            {
                "front_matter": self._front_matter(),
                "sections": [
                    {
                        "section_id": "sec1",
                        "title": "A focused narrative",
                        "paragraphs": [
                            {
                                "markdown": "A concise evidence-led discussion may choose its own paragraph pattern [@P001].",
                                "evidence_ids": ["P001-E01"],
                            }
                        ],
                    }
                ],
            },
        )
        result = run(BLUEPRINT_VALIDATOR, "--review-root", self.root, "--project-id", self.project_id)
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        result = run(DRAFT_VALIDATOR, "--review-root", self.root, "--project-id", self.project_id)
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        report = json.loads(
            (self.project / "02_section_drafting" / "section_draft_validation.json").read_text(
                encoding="utf-8"
            )
        )
        self.assertFalse(
            any(
                issue.startswith("draft_word_volume_below_80_percent_of_blueprint_target:")
                for issue in report["blocking_issues"]
            )
        )
        self.assertFalse(
            any("paragraph_type" in issue for issue in report["blocking_issues"])
        )
        report = json.loads(
            (self.project / "02_section_drafting" / "section_draft_validation.json").read_text(
                encoding="utf-8"
            )
        )
        self.assertTrue(any("planning estimate" in warning for warning in report["warnings"]))
        self.assertTrue(
            any(
                warning.startswith("draft_word_volume_below_80_percent_of_blueprint_target:")
                for warning in report["warnings"]
            )
        )

    def test_duplicate_template_paragraphs_are_blocked(self) -> None:
        write_json(
            self.project / "01_matrix_outline" / "section_blueprint.json",
            {
                "coverage_contract": {"dimensions": []},
                "sections": [{"section_id": "sec1", "title": "Review"}],
            },
        )
        repeated = (
            "This deliberately repeated review paragraph contains enough words to represent template padding. "
            "It makes a broad comparison without adding paper-specific chemical evidence, scope boundaries, "
            "reaction conditions, selectivity values, mechanistic distinctions, or meaningful interpretation "
            "for the reader, and therefore should never be counted twice in a finished manuscript"
        )
        write_json(
            self.project / "02_section_drafting" / "section_drafts.json",
            {
                "front_matter": self._front_matter(),
                "sections": [
                    {
                        "section_id": "sec1",
                        "title": "Review",
                        "paragraphs": [
                            {"paragraph_id": "sec1-p1", "markdown": repeated + " [@P001].", "evidence_ids": ["P001-E01"]},
                            {"paragraph_id": "sec1-p2", "markdown": repeated + " [@P002].", "evidence_ids": ["P002-E01"]},
                        ],
                    }
                ],
            },
        )
        result = run(DRAFT_VALIDATOR, "--review-root", self.root, "--project-id", self.project_id)
        self.assertNotEqual(result.returncode, 0)
        self.assertIn("duplicated long paragraph", result.stdout)

    def test_matrix_boilerplate_across_papers_is_blocked(self) -> None:
        repeated = (
            "The experimental procedures and characterization data fully support the proposed pathway, "
            "while systematic optimization establishes broad scope and practical utility for this method."
        )
        rows = [
            {
                "paper_id": f"P{i:03d}",
                "title": f"Distinct paper {i}",
                "role_after_reading": "background",
                "main_content": repeated,
            }
            for i in range(1, 6)
        ]
        write_json(
            self.project / "01_matrix_outline" / "literature_matrix.json", {"papers": rows}
        )
        result = run(MATRIX_VALIDATOR, "--review-root", self.root, "--project-id", self.project_id)
        self.assertNotEqual(result.returncode, 0)
        self.assertIn("boilerplate sentence repeated", result.stdout)

    def test_validation_report_becomes_stale_when_its_input_changes(self) -> None:
        result = run(
            MATRIX_VALIDATOR,
            "--review-root",
            self.root,
            "--project-id",
            self.project_id,
        )
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        matrix_path = self.project / "01_matrix_outline" / "literature_matrix.json"
        matrix_path.write_text(matrix_path.read_text(encoding="utf-8") + "\n", encoding="utf-8")
        spec = importlib.util.spec_from_file_location("status_stale_receipt", PROJECT_STATUS)
        module = importlib.util.module_from_spec(spec)
        assert spec.loader is not None
        spec.loader.exec_module(module)
        stage = next(item for item in module.STAGES if item["id"] == "matrix_outline")
        status = module.stage_status(self.project, stage)
        self.assertTrue(
            any("validation_input_stale" in issue for issue in status["semantic_issues"]),
            status["semantic_issues"],
        )

    def test_semantic_removed_verdict_requires_the_span_to_be_absent(self) -> None:
        spec = importlib.util.spec_from_file_location("audit_removed_contract", AUDIT)
        module = importlib.util.module_from_spec(spec)
        assert spec.loader is not None
        spec.loader.exec_module(module)
        manuscript = "This unsupported operational claim remains in the manuscript unchanged."
        audit_path = self.project / "05_final_audit" / "semantic_audit.json"
        write_json(
            audit_path,
            {
                "model": "fixture-reviewer",
                "manuscript_sha256": hashlib.sha256(manuscript.encode()).hexdigest(),
                "checks": [
                    {
                        "queue_id": "sec1-p1-a1",
                        "section_id": "sec1",
                        "paragraph_id": "sec1-p1",
                        "text_span": manuscript,
                        "verdict": "removed",
                        "comment": "The source does not report this operational condition, so the passage was deleted.",
                    }
                ],
                "unresolved_blockers": [],
            },
        )
        _, blockers = module.semantic_audit_status(
            audit_path,
            required=True,
            review_root=self.root,
            known_paper_ids=set(),
            evidence_owners={},
            evidence_records={},
            paragraphs_by_id={
                "sec1-p1": {"section_id": "sec1", "markdown": manuscript, "cited_paper_ids": [], "evidence_ids": []}
            },
            required_queue_ids={"sec1-p1-a1"},
            queue_items_by_id={
                "sec1-p1-a1": {"section_id": "sec1", "paragraph_id": "sec1-p1", "text_span": manuscript}
            },
            manuscript_sha256=hashlib.sha256(manuscript.encode()).hexdigest(),
            manuscript_text=manuscript,
        )
        self.assertTrue(any("marked removed" in issue for issue in blockers), blockers)

    def test_strict_reference_check_compares_title_authors_and_year_to_pdf(self) -> None:
        import fitz

        pdf = self.root / "review-library" / "sources" / "P001.pdf"
        pdf.parent.mkdir(parents=True, exist_ok=True)
        with fitz.open() as document:
            page = document.new_page()
            page.insert_text(
                (72, 90),
                "A Correct Catalytic Study\nAda Chemist and Lin Researcher\nPublished 2024\n10.1000/correct",
                fontsize=11,
            )
            document.save(pdf)
        metadata_path = self.root / "review-library" / "metadata" / "papers" / "P001.metadata.json"
        metadata = json.loads(metadata_path.read_text(encoding="utf-8"))
        metadata.update(
            {
                "title": {"value": "An Unrelated Fabricated Article"},
                "authors": {"value": ["Wrong, W."]},
                "year": {"value": 2021},
                "doi": {"value": "10.1000/correct"},
                "source_paths": {"pdf": str(pdf)},
            }
        )
        write_json(metadata_path, metadata)
        spec = importlib.util.spec_from_file_location("audit_reference_front", AUDIT)
        module = importlib.util.module_from_spec(spec)
        assert spec.loader is not None
        spec.loader.exec_module(module)
        conflicts = module.reference_metadata_source_conflicts(
            self.root,
            ["P001"],
            {"P001": {}},
            strict_front_matter=True,
        )
        issue_types = {item.get("issue") for item in conflicts}
        self.assertIn("title_not_supported_by_source_front_matter", issue_types)
        self.assertIn("authors_not_supported_by_source_front_matter", issue_types)
        self.assertIn("year_not_supported_by_source_front_matter", issue_types)

    def test_page_inspection_requires_hash_bound_page_specific_observations(self) -> None:
        spec = importlib.util.spec_from_file_location("render_inspection_contract", RENDER_DOCX)
        module = importlib.util.module_from_spec(spec)
        assert spec.loader is not None
        spec.loader.exec_module(module)
        page1 = self.project / "page-1.png"
        page2 = self.project / "page-2.png"
        page1.parent.mkdir(parents=True, exist_ok=True)
        page1.write_bytes(b"page one")
        page2.write_bytes(b"page two")
        inspection = self.project / "inspection.json"
        write_json(
            inspection,
            {
                "pages": [
                    {
                        "page_number": 1,
                        "page_sha256": hashlib.sha256(page1.read_bytes()).hexdigest(),
                        "verdict": "passed",
                        "observation": "The opening page shows the complete title and abstract with balanced margins and no clipping.",
                    }
                ]
            },
        )
        _, issues = module.load_page_inspections(inspection, [str(page1), str(page2)])
        self.assertTrue(any("uninspected pages" in issue for issue in issues), issues)


if __name__ == "__main__":
    unittest.main()
